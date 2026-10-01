from __future__ import annotations

import argparse
import json
import re
import sys
import time
import unicodedata
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from html import unescape
from pathlib import Path
from typing import Any
from urllib.parse import urlencode
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


BASE_URL = "https://ephemerajpp.com"
API_BASE = f"{BASE_URL}/wp-json/wp/v2"
SITE_NAME = "Ephemera"
CACHE_DIR = Path("Auditoria_Ephemera/cache")
POSTS_FULL_PATH = CACHE_DIR / "posts_full_content.json"
PAGES_PATH = CACHE_DIR / "pages.json"
CATEGORIES_PATH = CACHE_DIR / "categories.json"
XLSX_PATH = Path("Auditoria_Ephemera/Auditoria_Qualidade_Ephemera.xlsx")
XLSX_CORRIGIDO_PATH = Path("Auditoria_Ephemera/Auditoria_Qualidade_Ephemera_corrigido.xlsx")

USER_AGENT = "Mozilla/5.0 Codex Ephemera text body audit"
EXPECTED_POST_TOTAL: int | None = None


def slugify_filename(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")
    return value or "WordPress"


def configure_site(base_url: str, out_dir: str | None = None, site_name: str | None = None, xlsx_path: str | None = None) -> None:
    global BASE_URL, API_BASE, SITE_NAME, CACHE_DIR, POSTS_FULL_PATH, PAGES_PATH, CATEGORIES_PATH
    global XLSX_PATH, XLSX_CORRIGIDO_PATH, USER_AGENT
    BASE_URL = base_url.rstrip("/")
    API_BASE = f"{BASE_URL}/wp-json/wp/v2"
    SITE_NAME = site_name or re.sub(r"^www\.", "", BASE_URL.split("//", 1)[-1].split("/", 1)[0]) or "WordPress"
    safe_name = slugify_filename(SITE_NAME)
    out = Path(out_dir) if out_dir else Path(f"Auditoria_{safe_name}")
    CACHE_DIR = out / "cache"
    POSTS_FULL_PATH = CACHE_DIR / "posts_full_content.json"
    PAGES_PATH = CACHE_DIR / "pages.json"
    CATEGORIES_PATH = CACHE_DIR / "categories.json"
    XLSX_PATH = Path(xlsx_path) if xlsx_path else out / f"Auditoria_Qualidade_{safe_name}.xlsx"
    XLSX_CORRIGIDO_PATH = out / f"Auditoria_Qualidade_{safe_name}_corrigido.xlsx"
    USER_AGENT = f"Mozilla/5.0 Codex {SITE_NAME} text body audit"

STOPWORDS = {
    "a", "as", "o", "os", "e", "de", "da", "das", "do", "dos", "em", "no", "na", "nos", "nas",
    "um", "uma", "uns", "umas", "para", "por", "com", "sem", "ao", "aos", "à", "às", "que", "se",
    "the", "of", "and", "in", "with", "from", "for", "les", "des", "del", "los", "las", "und", "der",
}

KNOWN_SITE_TERMS = {
    "ephemera", "jpp", "wordpress", "tablepress", "pdf", "jpg", "jpeg", "png", "html", "http", "https",
    "www", "blog", "online", "facebook", "twitter", "youtube", "flickr", "instagram", "isbn",
}

CURATED_TYPOS = {
    "aacademia": "academia",
    "academco": "académico",
    "aentrada": "entrada",
    "afriique": "afrique",
    "agadecimentos": "agradecimentos",
    "agardecimento": "agradecimento",
    "agarelas": "aguarelas",
    "aimeirim": "almeirim",
    "aindicato": "sindicato",
    "aliianca": "aliança",
    "almeiida": "almeida",
    "amesterdao": "amesterdão",
    "amundial": "mundial",
    "andalazia": "andaluzia",
    "antiaeria": "antiaérea",
    "aolidariedade": "solidariedade",
    "aprsentado": "apresentado",
    "arquitrectura": "arquitetura",
    "artificsal": "artificial",
    "assasino": "assassino",
    "ateleier": "atelier",
    "aurarquicas": "autárquicas",
    "autolocantes": "autocolantes",
    "auutarquicas": "autárquicas",
    "barcelon": "barcelona",
    "barreeiro": "barreiro",
    "bassassinato": "assassinato",
    "bibioteca": "biblioteca",
    "bibllioteca": "biblioteca",
    "bnadeiras": "bandeiras",
    "boletiim": "boletim",
    "bolletino": "bollettino",
    "camapanha": "campanha",
    "campesisnos": "campesinos",
    "catalinha": "catalunha",
    "ccoperativa": "cooperativa",
    "centrode": "centro de",
    "cidadaoos": "cidadãos",
    "comitees": "comités",
    "comitesw": "comités",
    "comunisya": "comunista",
    "connflitos": "conflitos",
    "consleho": "conselho",
    "coooperativa": "cooperativa",
    "cooperatova": "cooperativa",
    "cooperatvo": "cooperativo",
    "cordeeeiros": "cordeiros",
    "correlegionarios": "correligionários",
    "ctibistas": "ativistas",
    "cultutal": "cultural",
    "curcular": "circular",
    "democratca": "democrática",
    "dezemrbo": "dezembro",
    "documentso": "documentos",
    "durng": "during",
    "eelicoes": "eleições",
    "eleicoies": "eleições",
    "eleiloes": "eleições",
    "eleitors": "eleitores",
    "elekcoes": "eleições",
    "emtradas": "entradas",
    "enblemas": "emblemas",
    "engenehiro": "engenheiro",
    "ensenble": "ensemble",
    "entradsa": "entradas",
    "ephemara": "ephemera",
    "espannha": "espanha",
    "esspanha": "espanha",
    "estudantss": "estudantes",
    "estuddos": "estudos",
}

FALSE_POSITIVE_TEXT_TERMS = {
    # Foreign words or names that are valid in context.
    "llau", "llegue", "posee", "supplemento", "drammaturgo", "counselling", "gwynn", "dall",
    "cabellos", "colecciones", "coleccionados", "sellos", "pissé", "chaisse", "mignonn", "homm",
    "persee", "metoo",
    # Historical Portuguese / pre-orthographic-agreement spellings commonly preserved in archive material.
    "reaccionários", "reaccionárias", "leccionou", "leccionadas", "inspeccionado", "coleccionou",
    "coleccionava", "coleccionismos", "seleccionadas", "seleccionada", "seleccionado", "seleccionem",
    "seleccionámos", "accionistas", "redireccionar", "diccionário", "occasionadas", "alli", "chamma",
    "impelles", "impellir", "janella", "soffrimento",
    # Abbreviations, fragments, or recurring extracted table artefacts.
    "ssica", "coord", "revolutionnair", "inff",
}

REPEATED_WORD_FALSE_POSITIVES = {
    "caricaturas", "pcp", "prp", "barreiro", "baleizão", "sipe", "preto", "callas", "coimbra",
    "storchio", "puna", "vicente", "social", "angola", "caça", "pcf", "tavares", "salazar",
    "corado", "cinfães", "almada", "alves", "margarida", "tenreiro", "robert", "eric", "gus",
    "carvalho", "girassol", "programa", "manifesto", "posters", "mural", "ephemera", "chega", "afasta",
}

REPEATED_WORD_KEEP = {"that", "about", "meia", "smart", "burst"}

SUGGESTION_OVERRIDES = {
    "reabiitação": "reabilitação",
    "reabiiitação": "reabilitação",
    "reabiiitacao": "reabilitação",
    "reabiiitação": "reabilitação",
    "reabiiitacão": "reabilitação",
    "reabiIitação": "reabilitação",
    "corrrer": "correr",
}


def strip_html(value: str | None) -> str:
    if not value:
        return ""
    value = re.sub(r"<script\b[^<]*(?:(?!</script>)<[^<]*)*</script>", " ", value, flags=re.I)
    value = re.sub(r"<style\b[^<]*(?:(?!</style>)<[^<]*)*</style>", " ", value, flags=re.I)
    value = re.sub(r"<[^>]+>", " ", value)
    value = re.sub(r"\[[a-zA-Z_][^\]]{0,180}\]", " ", value)
    return re.sub(r"\s+", " ", unescape(value)).strip()


def norm_word(value: str) -> str:
    value = unicodedata.normalize("NFKD", value.lower())
    return "".join(ch for ch in value if not unicodedata.combining(ch))


def edit_distance(a: str, b: str, max_distance: int = 2) -> int:
    a = norm_word(a)
    b = norm_word(b)
    if abs(len(a) - len(b)) > max_distance:
        return max_distance + 1
    previous = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        current = [i]
        row_min = i
        for j, cb in enumerate(b, 1):
            value = min(previous[j] + 1, current[-1] + 1, previous[j - 1] + (ca != cb))
            current.append(value)
            row_min = min(row_min, value)
        if row_min > max_distance:
            return max_distance + 1
        previous = current
    return previous[-1]


def fetch_json(url: str, retries: int = 3) -> tuple[Any, dict[str, str]]:
    req = Request(url, headers={"User-Agent": USER_AGENT, "Accept": "application/json"})
    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            with urlopen(req, timeout=90) as response:
                data = json.loads(response.read().decode("utf-8"))
                headers = {k.lower(): v for k, v in response.headers.items()}
                return data, headers
        except (HTTPError, URLError, TimeoutError) as exc:
            last_error = exc
            if attempt == retries:
                break
            time.sleep(2 * attempt)
    raise RuntimeError(f"API request failed after {retries} attempts: {url}") from last_error


def api_url(endpoint: str, **params: Any) -> str:
    return f"{API_BASE}/{endpoint}?{urlencode(params)}"


def fetch_posts_full_content() -> list[dict[str, Any]]:
    global EXPECTED_POST_TOTAL
    if POSTS_FULL_PATH.exists():
        posts = json.loads(POSTS_FULL_PATH.read_text(encoding="utf-8"))
        if posts and all("content" in post for post in posts[:10]):
            first_url = api_url("posts", per_page=25, page=1, _fields="id")
            _, headers = fetch_json(first_url)
            total_posts = int(headers.get("x-wp-total", str(len(posts))))
            EXPECTED_POST_TOTAL = total_posts
            if len(posts) >= total_posts:
                return posts
            if len(posts) >= total_posts - 25:
                print(f"Using near-complete post-content cache: {len(posts)}/{total_posts}", flush=True)
                return posts
            print(f"Resuming partial post-content cache: {len(posts)}/{total_posts}", flush=True)

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    fields = "id,date,modified,slug,link,title,excerpt,content,status,type"
    per_page = 25
    first_url = api_url("posts", per_page=per_page, page=1, _fields=fields)
    first_page, headers = fetch_json(first_url)
    total_pages = int(headers.get("x-wp-totalpages", "1"))
    EXPECTED_POST_TOTAL = int(headers.get("x-wp-total", "0") or 0)
    cached_posts: list[dict[str, Any]] = []
    if POSTS_FULL_PATH.exists():
        cached_posts = json.loads(POSTS_FULL_PATH.read_text(encoding="utf-8"))
    seen_ids = {post.get("id") for post in cached_posts}
    posts = list(cached_posts)
    for post in first_page:
        if post.get("id") not in seen_ids:
            posts.append(post)
            seen_ids.add(post.get("id"))
    POSTS_FULL_PATH.write_text(json.dumps(posts, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Fetched page 1/{total_pages}: {len(posts)} posts", flush=True)

    start_page = max(2, len(posts) // per_page + 1)

    def fetch_page(page_number: int) -> tuple[int, list[dict[str, Any]]]:
        url = api_url("posts", per_page=per_page, page=page_number, _fields=fields)
        data, _ = fetch_json(url)
        return page_number, data

    for batch_start in range(start_page, total_pages + 1, 30):
        batch_end = min(total_pages, batch_start + 29)
        pages = list(range(batch_start, batch_end + 1))
        fetched: list[tuple[int, list[dict[str, Any]]]] = []
        with ThreadPoolExecutor(max_workers=6) as executor:
            futures = {executor.submit(fetch_page, page): page for page in pages}
            for future in as_completed(futures):
                page = futures[future]
                try:
                    fetched.append(future.result())
                except Exception as exc:
                    print(f"Page {page} failed: {exc}", flush=True)
        for page, data in sorted(fetched, key=lambda x: x[0]):
            for post in data:
                if post.get("id") not in seen_ids:
                    posts.append(post)
                    seen_ids.add(post.get("id"))
        POSTS_FULL_PATH.write_text(json.dumps(posts, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"Fetched through page {batch_end}/{total_pages}: {len(posts)} posts", flush=True)
        time.sleep(0.2)

    POSTS_FULL_PATH.write_text(json.dumps(posts, ensure_ascii=False, indent=2), encoding="utf-8")
    return posts


def load_spellcheckers() -> dict[str, Any]:
    for vendor in [CACHE_DIR.parent / "vendor", Path("Auditoria_Ephemera/vendor")]:
        if vendor.exists() and str(vendor) not in sys.path:
            sys.path.insert(0, str(vendor))
    from spellchecker import SpellChecker  # type: ignore

    checkers: dict[str, Any] = {}
    for language in ["pt", "en", "es", "it", "fr", "de", "nl"]:
        try:
            checker = SpellChecker(language=language)
            checker.word_frequency.load_words(KNOWN_SITE_TERMS)
            checkers[language] = checker
        except Exception:
            pass
    return checkers


def add_domain_words(checkers: dict[str, Any], categories: list[dict[str, Any]], items: list[dict[str, Any]]) -> None:
    words: set[str] = set(KNOWN_SITE_TERMS)
    for category in categories:
        words.update(w.lower() for w in re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ]{4,}", strip_html(category.get("name"))))
    for item in items[:5000]:
        words.update(w.lower() for w in re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ]{4,}", strip_html((item.get("title") or {}).get("rendered"))))
    for checker in checkers.values():
        checker.word_frequency.load_words(words)


def raw_words(value: str) -> list[tuple[str, int]]:
    return [(m.group(0), m.start()) for m in re.finditer(r"[A-Za-zÀ-ÖØ-öø-ÿ]{4,}", value or "")]


def paragraph_language(paragraph: str, checkers: dict[str, Any]) -> str:
    counts: Counter[str] = Counter()
    words = [w.lower() for w, _ in raw_words(paragraph) if len(w) >= 5 and w.lower() not in STOPWORDS]
    for word in words[:80]:
        for lang, checker in checkers.items():
            if word in checker:
                counts[lang] += 1
    if not counts:
        return "indefinido"
    lang, count = counts.most_common(1)[0]
    total = sum(counts.values())
    if count >= 5 and count / max(total, 1) >= 0.45:
        return lang
    return "misto"


def is_known_in_any(word: str, checkers: dict[str, Any]) -> bool:
    return any(word in checker for checker in checkers.values())


def collapse_repeated_once(word: str) -> set[str]:
    variants = set()
    for match in re.finditer(r"([a-zà-ÿ])\1+", word):
        start, end = match.span()
        variants.add(word[:start] + match.group(1) + word[end:])
        variants.add(word[:start] + match.group(1) * 2 + word[end:])
    return variants


def likely_spelling_issue(raw: str, language: str, checkers: dict[str, Any]) -> tuple[str, str] | None:
    word = raw.lower()
    normalized = norm_word(word)
    if len(word) < 5 or word in STOPWORDS or word in KNOWN_SITE_TERMS:
        return None
    if raw.isupper() and len(raw) <= 8:
        return None
    if re.fullmatch(r"[ivxlcdmIVXLCDM]+", raw):
        return None
    if any(ch.isdigit() for ch in raw):
        return None
    if raw[0].isupper():
        return None
    if normalized in CURATED_TYPOS:
        return ("erro conhecido recuperado", CURATED_TYPOS[normalized])
    if is_known_in_any(word, checkers):
        return None

    for variant in collapse_repeated_once(word):
        if len(variant) >= 4 and is_known_in_any(variant, checkers):
            return ("letra repetida indevidamente", variant)

    candidate_langs = [language] if language in checkers else ["pt"]
    if language in {"misto", "indefinido"}:
        candidate_langs = ["pt", "en", "es", "it", "fr"]
    for lang in candidate_langs:
        checker = checkers.get(lang)
        if not checker:
            continue
        suggestion = checker.correction(word)
        if not suggestion or suggestion == word:
            continue
        distance = edit_distance(word, suggestion, 2)
        if distance <= 2 and suggestion[0] == word[0]:
            return (f"erro provável por dicionário ({lang})", suggestion)
    return None


def context_excerpt(text: str, index: int, width: int = 95) -> str:
    start = max(0, index - width)
    end = min(len(text), index + width)
    excerpt = text[start:end].strip()
    return re.sub(r"\s+", " ", excerpt)


def text_fields(item: dict[str, Any], content_type: str) -> list[tuple[str, str]]:
    title = strip_html((item.get("title") or {}).get("rendered"))
    excerpt = strip_html((item.get("excerpt") or {}).get("rendered"))
    content = strip_html((item.get("content") or {}).get("rendered"))
    fields = [("Título", title), ("Excerto", excerpt)]
    if content:
        fields.append(("Corpo", content))
    elif content_type == "post":
        fields.append(("Corpo", ""))
    return fields


def analyze_items(items: list[tuple[str, dict[str, Any]]], checkers: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, str]] = set()
    curated_re = re.compile(r"\b(" + "|".join(re.escape(k) for k in sorted(CURATED_TYPOS, key=len, reverse=True)) + r")\b", re.I)

    for content_type, item in items:
        url = item.get("link") or ""
        title = strip_html((item.get("title") or {}).get("rendered"))
        for field, text in text_fields(item, content_type):
            if not text:
                continue
            paragraphs = [p.strip() for p in re.split(r"(?<=[.!?])\s+|\n+", text) if p.strip()]
            for paragraph in paragraphs:
                language = "não classificado"

                for match in re.finditer(r"\b([A-Za-zÀ-ÖØ-öø-ÿ]{3,})\s+\1\b", paragraph, flags=re.I):
                    repeated = match.group(1)
                    if repeated.lower() in {"que", "the", "and", "para"}:
                        continue
                    key = (url, field, "palavra repetida consecutivamente", repeated.lower())
                    if key not in seen:
                        seen.add(key)
                        rows.append({
                            "Tipo conteúdo": content_type,
                            "Campo": field,
                            "Tipo de suspeita": "Palavra repetida consecutivamente",
                            "Termo": repeated,
                            "Sugestão": f"remover repetição de '{repeated}'",
                            "Idioma provável": language,
                            "URL": url,
                            "Título": title,
                            "Contexto": context_excerpt(paragraph, match.start()),
                            "Confiança": "Alta",
                        })

                if re.search(r"[�]|Ã[\x80-\xbf]", paragraph):
                    key = (url, field, "possível problema de codificação", paragraph[:50])
                    if key not in seen:
                        seen.add(key)
                        rows.append({
                            "Tipo conteúdo": content_type,
                            "Campo": field,
                            "Tipo de suspeita": "Possível problema de codificação",
                            "Termo": "",
                            "Sugestão": "rever caracteres acentuados/encoding",
                            "Idioma provável": language,
                            "URL": url,
                            "Título": title,
                            "Contexto": context_excerpt(paragraph, 0),
                            "Confiança": "Alta",
                        })

                for match in curated_re.finditer(paragraph):
                    raw = match.group(1)
                    suggestion = CURATED_TYPOS[norm_word(raw)]
                    key = (url, field, "erro conhecido recuperado", raw.lower())
                    if key in seen:
                        continue
                    seen.add(key)
                    rows.append({
                        "Tipo conteúdo": content_type,
                        "Campo": field,
                        "Tipo de suspeita": "Erro conhecido recuperado",
                        "Termo": raw,
                        "Sugestão": suggestion,
                        "Idioma provável": language,
                        "URL": url,
                        "Título": title,
                        "Contexto": context_excerpt(paragraph, match.start()),
                        "Confiança": "Alta",
                    })

                for raw, index in raw_words(paragraph):
                    word = raw.lower()
                    if not re.search(r"([a-zà-ÿ])\1{1,}", word) and not re.search(r"([a-zà-ÿ])\1{2,}", word):
                        continue
                    if raw.isupper() or raw[0].isupper() or re.fullmatch(r"[ivxlcdmIVXLCDM]+", raw):
                        continue
                    if is_known_in_any(word, checkers):
                        continue
                    suggestion = ""
                    for variant in collapse_repeated_once(word):
                        if is_known_in_any(variant, checkers):
                            suggestion = variant
                            break
                    if not suggestion:
                        continue
                    issue_type = "Letra repetida indevidamente"
                    key = (url, field, issue_type, raw.lower())
                    if key in seen:
                        continue
                    seen.add(key)
                    rows.append({
                        "Tipo conteúdo": content_type,
                        "Campo": field,
                        "Tipo de suspeita": issue_type,
                        "Termo": raw,
                        "Sugestão": suggestion,
                        "Idioma provável": language,
                        "URL": url,
                        "Título": title,
                        "Contexto": context_excerpt(paragraph, index),
                        "Confiança": "Alta",
                    })
                    if len(rows) >= 1000:
                        break
                if len(rows) >= 1000:
                    break
            if len(rows) >= 1000:
                break
        if len(rows) >= 1000:
            break

    cleaned_rows: list[dict[str, Any]] = []
    cleaned_seen: set[tuple[str, str, str, str, str]] = set()
    for row in rows:
        term = str(row.get("Termo") or "")
        term_key = term.lower()
        issue_type = str(row.get("Tipo de suspeita") or "")
        if issue_type == "Palavra repetida consecutivamente":
            if term_key not in REPEATED_WORD_KEEP:
                continue
            # Keep only exact lower-case repetitions in prose-like contexts.
            if not re.search(rf"\b{re.escape(term)}\s+{re.escape(term)}\b", row.get("Contexto") or ""):
                continue
        if issue_type == "Letra repetida indevidamente":
            if term_key in {t.lower() for t in FALSE_POSITIVE_TEXT_TERMS}:
                continue
            if term in SUGGESTION_OVERRIDES:
                row["Sugestão"] = SUGGESTION_OVERRIDES[term]
            elif term_key in SUGGESTION_OVERRIDES:
                row["Sugestão"] = SUGGESTION_OVERRIDES[term_key]
        key = (
            str(row.get("URL") or ""),
            str(row.get("Campo") or ""),
            issue_type,
            term_key,
            str(row.get("Contexto") or "")[:80],
        )
        if key in cleaned_seen:
            continue
        cleaned_seen.add(key)
        cleaned_rows.append(row)

    rows = cleaned_rows
    summary = []
    by_field = Counter(row["Campo"] for row in rows)
    by_type = Counter(row["Tipo de suspeita"] for row in rows)
    by_confidence = Counter(row["Confiança"] for row in rows)
    for key, value in {
        "Itens analisados": len(items),
        "Suspeitas textuais encontradas": len(rows),
        "Posts com corpo recolhido": sum(1 for typ, item in items if typ == "post" and (item.get("content") or {}).get("rendered")),
        "Posts não recolhidos por erro API": max((EXPECTED_POST_TOTAL or 0) - sum(1 for typ, _ in items if typ == "post"), 0),
        "Páginas com corpo analisado": sum(1 for typ, item in items if typ == "page" and (item.get("content") or {}).get("rendered")),
    }.items():
        summary.append({"Indicador": key, "Valor": value})
    for field, value in by_field.items():
        summary.append({"Indicador": f"Campo: {field}", "Valor": value})
    for issue_type, value in by_type.items():
        summary.append({"Indicador": f"Tipo: {issue_type}", "Valor": value})
    for confidence, value in by_confidence.items():
        summary.append({"Indicador": f"Confiança: {confidence}", "Valor": value})
    return rows, summary


def write_sheet(wb: Any, name: str, rows: list[dict[str, Any]]) -> None:
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    if not rows:
        ws.append(["Sem resultados"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {
        "A": 14, "B": 12, "C": 34, "D": 20, "E": 24, "F": 14, "G": 52, "H": 48, "I": 80, "J": 12,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default=BASE_URL, help="URL base do site WordPress, por exemplo https://exemplo.com")
    parser.add_argument("--site-name", default=None, help="Nome curto usado no relatório e nos nomes dos ficheiros.")
    parser.add_argument("--out-dir", default=None, help="Pasta da auditoria criada pelo script principal.")
    parser.add_argument("--xlsx-path", default=None, help="Caminho opcional para o Excel onde serão adicionadas as folhas textuais.")
    args = parser.parse_args()

    configure_site(args.base_url, out_dir=args.out_dir, site_name=args.site_name, xlsx_path=args.xlsx_path)

    posts = fetch_posts_full_content()
    pages = json.loads(PAGES_PATH.read_text(encoding="utf-8"))
    categories = json.loads(CATEGORIES_PATH.read_text(encoding="utf-8"))
    items = [("post", item) for item in posts] + [("page", item) for item in pages]

    checkers = load_spellcheckers()
    add_domain_words(checkers, categories, posts + pages)
    rows, summary = analyze_items(items, checkers)

    out_json = CACHE_DIR / "text_body_audit_results.json"
    out_json.write_text(json.dumps({"summary": summary, "rows": rows}, ensure_ascii=False, indent=2), encoding="utf-8")

    wb = load_workbook(XLSX_PATH)
    write_sheet(wb, "Auditoria Texto", rows)
    write_sheet(wb, "Resumo Texto", summary)
    wb.save(XLSX_PATH)
    wb.save(XLSX_CORRIGIDO_PATH)
    print(f"Textual audit rows: {len(rows)}")
    print(Counter(row['Campo'] for row in rows))
    print(Counter(row['Tipo de suspeita'] for row in rows))


if __name__ == "__main__":
    main()
